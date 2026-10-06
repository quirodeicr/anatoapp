# -*- coding: utf-8 -*-
"""
Extrae las fotos y los textos de "Mat 3 Props.xlsx" (ejercicios con props:
bola, ligas, pesa, círculo, roller) a privado/mat3/, que NO se versiona: el
repositorio es público y las fotos muestran personas (y las observaciones,
datos personales).

  python importar_mat3.py "ruta/Mat 3 Props.xlsx"

Cada hoja tiene la ficha (nombre, posición, principio, objetivo, un renglón por
prop y observaciones) y fotos ancladas en el renglón del prop que muestran.
Las caras de quienes miran la clase se pixelan (CARAS); quien hace el ejercicio
queda visible porque es la demostración. Las fotos se achican a JPEG.
Salida: privado/mat3/<hoja>/NN.jpg y privado/mat3/mat3.json.
"""
import io, os, re, sys, json
import openpyxl
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(BASE, "privado", "mat3")
ANCHO = 720          # ancho máximo de cada foto
CALIDAD = 72

# hoja → ejercicio de la app (solo las hojas ya revisadas foto por foto)
EJERCICIO = {"Roll Up 1": "mat1-e03", "SWAN": "mat1-e15"}

# caras de quienes no hacen el ejercicio: (hoja, n.º de foto) → cajas [x0, y0, x1, y1]
# en fracciones de la imagen (las fotos se numeran por fila y columna de anclaje)
CARAS = {
    ("Roll Up 1", 3): [[0.55, 0, 0.68, 0.17], [0.77, 0, 0.9, 0.17]],
    ("Roll Up 1", 6): [[0.53, 0.03, 0.65, 0.2], [0.77, 0.07, 0.89, 0.24]],
    ("Roll Up 1", 7): [[0.73, 0, 0.85, 0.16]],
    ("Roll Up 1", 8): [[0.42, 0, 0.6, 0.07], [0.76, 0, 0.92, 0.06]],
    ("Roll Up 1", 10): [[0.2, 0.02, 0.35, 0.25], [0.45, 0, 0.56, 0.09], [0.7, 0, 0.82, 0.17]],
}

def pixelar(im, caja):
    w, h = im.size
    x0, y0, x1, y1 = int(caja[0] * w), int(caja[1] * h), int(caja[2] * w), int(caja[3] * h)
    zona = im.crop((x0, y0, x1, y1))
    bloque = max(8, w // 60)
    chica = zona.resize((max(1, (x1 - x0) // bloque), max(1, (y1 - y0) // bloque)), Image.BILINEAR)
    im.paste(chica.resize(zona.size, Image.NEAREST), (x0, y0))

def carpeta(nombre):
    return re.sub(r"[^a-z0-9]+", "-", nombre.lower()).strip("-")

def texto(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

def main(ruta):
    wb = openpyxl.load_workbook(ruta)
    hojas = []
    for ws in wb.worksheets:
        imgs = sorted(ws._images, key=lambda i: (i.anchor._from.row, i.anchor._from.col, i.anchor._from.rowOff, i.anchor._from.colOff))
        if not imgs:
            continue
        etiquetas = {r: texto(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1)}
        valor = {etiquetas[r]: texto(ws.cell(r, 2).value) for r in etiquetas if etiquetas[r]}
        props = {k: valor.get(k, "") for k in ("Bola", "Liga Cerrada", "Liga Abierta", "Pesa", "Círculo", "Roler")}
        dest = os.path.join(SALIDA, carpeta(ws.title))
        os.makedirs(dest, exist_ok=True)
        fotos = []
        for n, img in enumerate(imgs, 1):
            im = Image.open(io.BytesIO(img._data())).convert("RGB")
            for caja in CARAS.get((ws.title, n), []):
                pixelar(im, caja)
            if im.width > ANCHO:
                im = im.resize((ANCHO, round(im.height * ANCHO / im.width)), Image.LANCZOS)
            archivo = f"{n:02d}.jpg"
            im.save(os.path.join(dest, archivo), "JPEG", quality=CALIDAD, optimize=True, progressive=True)
            prop = etiquetas.get(img.anchor._from.row + 1, "")
            fotos.append({"archivo": f"{carpeta(ws.title)}/{archivo}", "prop": "Roller" if prop == "Roler" else prop, "w": im.width, "h": im.height})
        hojas.append({"hoja": ws.title, "ejercicio": EJERCICIO.get(ws.title), "nombre": valor.get("Nombre del Ejercicio", ""),
                      "posicion": valor.get("Posición", ""), "principio": valor.get("Principio del Movimiento", ""),
                      "objetivo": valor.get("Objetivo", ""), "props": props, "observaciones": valor.get("Observaciones", ""), "fotos": fotos})
    json.dump(hojas, open(os.path.join(SALIDA, "mat3.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    peso = sum(os.path.getsize(os.path.join(SALIDA, f["archivo"])) for h in hojas for f in h["fotos"])
    print(len(hojas), "hojas,", sum(len(h["fotos"]) for h in hojas), f"fotos ({peso / 1024 / 1024:.1f} MB) →", SALIDA)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
